"""Read the material table of the client's `tank calculations.xlsx` ('links' sheet, J12:R45) and write it to
app/seed/data/tank_materials.json, which is what seed_tank_materials loads into the database.

Run only when the client sends a new sheet (from backend/):
    python -m app.seed.extract_tank_sheet "../tank calculations.xlsx"
"""
import json
import os
import sys

import openpyxl

OUT = os.path.join(os.path.dirname(__file__), "data", "tank_materials.json")
FIRST_ROW, LAST_ROW = 13, 45          # row 12 is the header


def _num(v):
    return None if v in (None, "", "NA") else float(v)


def read_materials(path: str):
    ws = openpyxl.load_workbook(path, data_only=True)["links"]
    out = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        spec, grade = ws.cell(r, 11).value, ws.cell(r, 12).value        # K, L
        if not spec:
            continue
        spec = str(spec).strip()
        grade = str(grade).strip() if grade not in (None, "") else None
        is_unknown = spec.lower() == "unknown"
        name = "Unknown Grade" if is_unknown else (f"{spec} Grade {grade}" if grade else spec)
        out.append({
            "name": name, "spec": spec, "grade": grade, "is_unknown": is_unknown,
            "yield_mpa": _num(ws.cell(r, 13).value),            # M
            "tensile_mpa": _num(ws.cell(r, 14).value),          # N
            "s_lower_mpa": _num(ws.cell(r, 15).value),          # O  allowable product stress, lower two courses
            "s_upper_mpa": _num(ws.cell(r, 16).value),          # P  upper courses
            "st_lower_mpa": _num(ws.cell(r, 17).value),         # Q  hydrostatic test stress, lower two courses
            "st_upper_mpa": _num(ws.cell(r, 18).value),         # R  upper courses
            "sheet_row": r,
        })
    return out


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "..", "..", "..", "tank calculations.xlsx")
    rows = read_materials(src)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=1, ensure_ascii=False)
    print(f"{len(rows)} materials written to {OUT}")
