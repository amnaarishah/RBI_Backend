"""The inspection import template, read strictly (own temporary database).

Covers the client's 51-E-01 file (the minimum of each inspection must be the real minimum of ALL its readings),
fewer than four readings, zone types, dates, zone spelling, duplicates, all-or-nothing, the downloadable template
(round trip) and CSV.

Run from backend/:  python -m unittest tests.test_inspection_template_import -v
"""
import io
import os
import sys
import unittest
import uuid
from datetime import datetime
from decimal import Decimal as D

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl  # noqa: E402
from sqlalchemy import create_engine, text  # noqa: E402
from sqlalchemy.engine import make_url  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app.core.config import settings  # noqa: E402
from app.core.database import Base  # noqa: E402
import app.models.equipment, app.models.inspection, app.models.cml  # noqa: E402,F401
from app.models.equipment import Equipment  # noqa: E402
from app.models.inspection import InspectionRecord  # noqa: E402
from app.models.cml import InspectionTML  # noqa: E402
from app.crud.inspection_import import import_inspections_from_excel  # noqa: E402
from app.crud import inspection_template as tpl  # noqa: E402

BASE_URL = make_url(settings.DATABASE_URL)
TEST_DB = "ams_test_" + uuid.uuid4().hex[:8]
HEADER = ["Equipment Tag", "Inspection Date", "Inspector Name", "TML Code", "TML Name", "TML Type", "Direction Labels",
          "Reading 1 (mm)", "Reading 2 (mm)", "Reading 3 (mm)", "Reading 4 (mm)", "Min Required Thickness (mm)", "Remarks"]
Dt = datetime


def row(tag, date, who, code, typ, r, labels="N, S, E, W", name=None, min_t=None, remarks=None):
    r = list(r) + [None] * (4 - len(r))
    return [tag, date, who, code, name or code, typ, labels, *r, min_t, remarks]


def make_xlsx(rows, header=HEADER, sheet="Inspection Data", extra_sheets=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    for name, data in extra_sheets:
        w2 = wb.create_sheet(name)
        for r in data:
            w2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv(rows, header=HEADER):
    def cell(v):
        if v is None:
            return ""
        return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else str(v)
    import csv
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\n")
    w.writerow(header)
    for r in rows:
        w.writerow([cell(v) for v in r])
    return out.getvalue().encode("utf-8")


class TemplateImport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        admin = create_engine(BASE_URL.set(database="postgres"), isolation_level="AUTOCOMMIT")
        with admin.connect() as c:
            c.execute(text(f'CREATE DATABASE "{TEST_DB}"'))
        admin.dispose()
        cls.engine = create_engine(BASE_URL.set(database=TEST_DB))
        Base.metadata.create_all(cls.engine)
        cls.Session = sessionmaker(bind=cls.engine)

    @classmethod
    def tearDownClass(cls):
        cls.engine.dispose()
        admin = create_engine(BASE_URL.set(database="postgres"), isolation_level="AUTOCOMMIT")
        with admin.connect() as c:
            c.execute(text(f'DROP DATABASE IF EXISTS "{TEST_DB}" WITH (FORCE)'))
        admin.dispose()

    def setUp(self):
        self.db = self.Session()
        # clean slate for every test
        for t in ("inspection_reading", "inspection_tml", "inspection_records", "equipment"):
            self.db.execute(text(f"DELETE FROM {t}"))
        self.db.commit()
        for tag, typ in (("51-E-01", "Pressure Vessel"), ("051-V-11", "Pressure Vessel"), ("00-E-12", "Heat Exchanger")):
            self.db.add(Equipment(equipment_id=tag, equipment_name=tag, equipment_type=typ, plant="P", unit="U",
                                  min_required_thickness=D("8.0"), dish_min_thickness=D("8.0"), tube_min_thickness=D("1.0"),
                                  bottom_min_thickness=D("4.0")))
        self.db.commit()

    def tearDown(self):
        self.db.close()

    def run_import(self, data, name="x.xlsx"):
        return import_inspections_from_excel(self.db, data, name)

    def records(self, tag):
        eq = self.db.query(Equipment).filter(Equipment.equipment_id == tag).one()
        return sorted(self.db.query(InspectionRecord).filter(InspectionRecord.equipment_id == eq.id).all(),
                      key=lambda r: r.inspection_date)

    # ------------------------------------------------------------------ the client's file
    def test_client_file_minimum_is_the_real_minimum_of_all_readings(self):
        rows = [
            row("51-E-01", Dt(2024, 1, 1), "John Doe", "S-1", "Shell", (10.01, 10.05, 10.06, 10.08)),
            row("51-E-01", Dt(2024, 1, 1), "John Doe", "S-2", "Shell", (10.04, 10.04, 10.00, 10.02)),
            row("51-E-01", Dt(2025, 1, 1), "John Doe", "S-1", "Shell", (9.98, 9.99, 9.96, 9.99)),
            row("51-E-01", Dt(2025, 1, 1), "John Doe", "S-2", "Shell", (9.963, 9.96, 9.95, 9.95)),
            row("51-E-01", Dt(2026, 1, 1), "Sarah Smith", "S-1", "Shell", (9.88, 9.89, 9.86, 9.88)),
            row("51-E-01", Dt(2026, 1, 1), "Sarah Smith", "S-2", "Shell", (9.86, 9.88, 9.88, 9.87)),
        ]
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(res["records_created"], 3)
        recs = self.records("51-E-01")
        self.assertEqual([r.min_measured_thickness for r in recs], [D("10.00"), D("9.95"), D("9.86")])
        # both zones and all four readings are stored (the old reader kept only Reading 4 of the first zone)
        for r in self.db.query(InspectionTML).all():
            self.assertEqual(r.tml_type, "Shell")
        self.assertEqual({z.tml_code for z in self.db.query(InspectionTML).all()}, {"S-1", "S-2"})
        s = res["summary"][0]
        self.assertEqual(len(s["zones"]), 2)
        self.assertEqual(len(s["zones"][0]["readings"]), 4)
        self.assertEqual(s["min_thickness"], "10.0000")
        self.assertIsNotNone(res["summary"][2]["remaining_life"])
        # every zone reports its own calculated result (baseline: none)
        self.assertTrue(all(z["remaining_life"] is None for z in res["summary"][0]["zones"]))
        z26 = {z["code"]: z for z in res["summary"][2]["zones"]}
        self.assertIsNotNone(z26["S-1"]["remaining_life"])
        self.assertIsNotNone(z26["S-1"]["corrosion_rate"])
        self.assertIn(z26["S-1"]["controlling_point"], (1, 2, 3, 4))

    # ------------------------------------------------------------------ fewer than four readings
    def test_three_two_and_one_reading_zones(self):
        rows = [
            row("00-E-12", Dt(2026, 1, 1), "A", "BTM-1", "Tank Bottom", (8.98, 4.98, 3.10)),
            row("00-E-12", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.9, 9.5)),
            row("00-E-12", Dt(2026, 1, 1), "A", "TUBE-1", "Tube", (2.01,), labels="Wall"),
        ]
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "success", res["errors"])
        rec = self.records("00-E-12")[0]
        self.assertEqual(rec.min_measured_thickness, D("2.01"))          # the tube is the lowest reading
        zones = {z["code"]: z for z in res["summary"][0]["zones"]}
        self.assertEqual(len(zones["BTM-1"]["readings"]), 3)
        self.assertEqual(zones["S-1"]["min_reading"], "9.5")

    def test_empty_and_not_measured_cells_are_never_zero(self):
        rows = [row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (10.2, "NA", "-", 10.0))]
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(self.records("51-E-01")[0].min_measured_thickness, D("10.0"))

    # ------------------------------------------------------------------ nothing is guessed
    def assert_rejected(self, rows, fragment, **kw):
        res = self.run_import(make_xlsx(rows, **kw))
        self.assertEqual(res["status"], "rejected", res)
        self.assertTrue(any(fragment in e for e in res["errors"]), res["errors"])
        self.assertEqual(self.db.query(InspectionRecord).count(), 0)       # nothing imported
        return res

    def test_missing_type_is_refused_not_defaulted_to_shell(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "T-1", None, (2.0,))], "TML Type: is empty")

    def test_invalid_type_is_refused(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "T-1", "Pipe", (2.0,))], "TML Type 'Pipe'")

    def test_unknown_equipment_is_refused(self):
        self.assert_rejected([row("99-X-99", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,))], "not in the Equipment Register")

    def test_leading_zero_makes_a_different_equipment(self):
        res = self.run_import(make_xlsx([row("51-V-11", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,))]))
        self.assertEqual(res["status"], "rejected")                     # 051-V-11 exists, 51-V-11 does not

    def test_bad_numbers_are_reported_with_row_and_column(self):
        rows = [row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", ("9,96", 0, "abc", 9.0))]
        res = self.assert_rejected(rows, "Row 2, Reading 1")
        text_ = " ".join(res["errors"])
        self.assertIn("Row 2, Reading 2", text_)
        self.assertIn("Row 2, Reading 3", text_)

    def test_row_with_no_reading_is_refused(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", ())], "no thickness reading")

    def test_one_wrong_row_blocks_the_whole_file(self):
        rows = [row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,)),
                row("51-E-01", Dt(2026, 1, 1), "A", "S-2", "Cone", (9.0,))]
        self.assert_rejected(rows, "TML Type 'Cone'")

    def test_dates(self):
        self.assert_rejected([row("51-E-01", "01/02/2026", "A", "S-1", "Shell", (9.0,))], "day/month or month/day")
        self.assert_rejected([row("51-E-01", "not a date", "A", "S-1", "Shell", (9.0,))], "is not a date")
        self.assert_rejected([row("51-E-01", Dt(2999, 1, 1), "A", "S-1", "Shell", (9.0,))], "in the future")
        ok = self.run_import(make_xlsx([row("51-E-01", "2026-03-04", "A", "S-1", "Shell", (9.0,)),
                                        row("51-E-01", "25/12/2025", "A", "S-1", "Shell", (9.1,)),
                                        row("51-E-01", "08-Aug-2025", "A", "S-1", "Shell", (9.2,))]))
        self.assertEqual(ok["status"], "success", ok["errors"])
        self.assertEqual([r.inspection_date.strftime("%Y-%m-%d") for r in self.records("51-E-01")],
                         ["2025-08-08", "2025-12-25", "2026-03-04"])

    def test_inspector_is_required(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), None, "S-1", "Shell", (9.0,))], "Inspector Name: is empty")

    def test_duplicate_zone_and_two_inspectors_in_one_inspection(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,)),
                              row("51-E-01", Dt(2026, 1, 1), "A", "S1", "Shell", (9.1,))], "is already in row 2")
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,)),
                              row("51-E-01", Dt(2026, 1, 1), "B", "S-2", "Shell", (9.1,))], "One inspection has one inspector")

    def test_labels_must_cover_the_readings(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0, 9.1, 9.2), labels="N, S")], "Direction Labels")

    def test_non_mm_unit_is_refused(self):
        header = HEADER + ["Unit"]
        rows = [row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,)) + ["inch"]]
        self.assert_rejected(rows, "must be in mm", header=header)

    def test_missing_required_column(self):
        header = [h for h in HEADER if h != "TML Type"]
        data = make_xlsx([[v for v, h in zip(row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0,)), HEADER) if h != "TML Type"]], header=header)
        with self.assertRaises(ValueError) as cm:
            self.run_import(data)
        self.assertIn("TML Type", str(cm.exception))

    def test_unknown_columns_are_reported_not_read(self):
        header = HEADER + ["min"]
        rows = [row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.0, 9.5)) + [1.0]]
        res = self.run_import(make_xlsx(rows, header=header))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertTrue(any("NOT read" in w and "min" in w for w in res["warnings"]), res["warnings"])
        self.assertEqual(self.records("51-E-01")[0].min_measured_thickness, D("9.0"))

    # ------------------------------------------------------------------ zones
    def test_same_zone_written_differently_keeps_its_history(self):
        self.run_import(make_xlsx([row("51-E-01", Dt(2025, 1, 1), "A", "S-1", "Shell", (10.0, 10.1))]))
        res = self.run_import(make_xlsx([row("51-E-01", Dt(2026, 1, 1), "A", "s1", "Shell", (9.8, 9.9))]))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(self.db.query(InspectionTML).count(), 1)         # one zone, not two
        self.assertTrue(any("matched to the existing zone 'S-1'" in w for w in res["warnings"]))
        self.assertIsNotNone(self.records("51-E-01")[-1].remaining_life)

    def test_new_zone_name_is_flagged(self):
        self.run_import(make_xlsx([row("51-E-01", Dt(2025, 1, 1), "A", "S-1", "Shell", (10.0,))]))
        res = self.run_import(make_xlsx([row("51-E-01", Dt(2026, 1, 1), "A", "Shell1", "Shell", (9.8,))]))
        self.assertEqual(res["status"], "success")
        self.assertTrue(any("NEW zone" in w and "S-1" in w for w in res["warnings"]), res["warnings"])

    def test_type_of_an_existing_zone_cannot_change(self):
        self.run_import(make_xlsx([row("51-E-01", Dt(2025, 1, 1), "A", "S-1", "Shell", (10.0,))]))
        res = self.run_import(make_xlsx([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Dish", (9.8,))]))
        self.assertEqual(res["status"], "rejected")
        self.assertTrue(any("already a Shell zone" in e for e in res["errors"]), res["errors"])

    def test_nozzle_needs_its_own_minimum(self):
        self.assert_rejected([row("51-E-01", Dt(2026, 1, 1), "A", "N-1", "Nozzle", (9.0,))], "Min Required Thickness: needed")
        ok = self.run_import(make_xlsx([row("51-E-01", Dt(2026, 1, 1), "A", "N-1", "Nozzle", (9.0,), min_t=6.5)]))
        self.assertEqual(ok["status"], "success", ok["errors"])
        self.assertEqual(ok["summary"][0]["zones"][0]["t_min"], "6.5")

    # ------------------------------------------------------------------ minimum thickness: register first, sheet fills a gap
    def equipment(self, tag):
        self.db.expire_all()
        return self.db.query(Equipment).filter(Equipment.equipment_id == tag).one()

    def test_register_has_none_and_the_sheet_has_one_it_is_used_and_saved_to_the_register(self):
        eq = self.equipment("051-V-11")
        eq.tube_min_thickness = None
        self.db.commit()
        rows = [row("051-V-11", Dt(2025, 1, 1), "A", "T-1", "Tube", (2.0,), min_t=1.5),
                row("051-V-11", Dt(2026, 1, 1), "A", "T-1", "Tube", (1.9,), min_t=1.5)]
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(self.equipment("051-V-11").tube_min_thickness, D("1.5"))
        self.assertTrue(any("saved to the Equipment Register" in w for w in res["warnings"]), res["warnings"])
        self.assertEqual(D(res["summary"][0]["zones"][0]["t_min"]), D("1.5"))
        zone = self.db.query(InspectionTML).filter(InspectionTML.tml_code == "T-1").one()
        self.assertEqual(zone.min_required_thickness, D("1.5"))
        self.assertIsNotNone(self.records("051-V-11")[-1].remaining_life)      # (1.9-1.5)/0.1 -> can be calculated now

    def test_register_and_sheet_both_have_one_the_register_wins(self):
        rows = [row("051-V-11", Dt(2026, 1, 1), "A", "T-1", "Tube", (2.0,), min_t=1.5)]      # register tube min = 1.0
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(self.equipment("051-V-11").tube_min_thickness, D("1.0"))          # unchanged
        self.assertEqual(D(res["summary"][0]["zones"][0]["t_min"]), D("1.0"))
        self.assertTrue(any("register always wins" in w for w in res["warnings"]), res["warnings"])

    def test_same_value_in_register_and_sheet_gives_no_warning(self):
        res = self.run_import(make_xlsx([row("051-V-11", Dt(2026, 1, 1), "A", "T-1", "Tube", (2.0,), min_t=1.0)]))
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertFalse(any("always wins" in w for w in res["warnings"]), res["warnings"])

    def test_no_minimum_anywhere_is_a_warning_not_an_error(self):
        eq = self.equipment("051-V-11")
        eq.tube_min_thickness = None
        self.db.commit()
        res = self.run_import(make_xlsx([row("051-V-11", Dt(2026, 1, 1), "A", "T-1", "Tube", (2.0,))]))
        self.assertEqual(res["status"], "success", res["errors"])           # saved
        self.assertEqual(res["records_created"], 1)
        self.assertTrue(any("neither in the Equipment Register nor in the file" in w for w in res["warnings"]), res["warnings"])
        self.assertIsNone(self.equipment("051-V-11").tube_min_thickness)    # nothing invented

    def test_two_different_sheet_minimums_for_the_same_register_field_are_refused(self):
        eq = self.equipment("051-V-11")
        eq.dish_min_thickness = None
        self.db.commit()
        rows = [row("051-V-11", Dt(2026, 1, 1), "A", "D-1", "Dish", (9.0,), min_t=8.0),
                row("051-V-11", Dt(2026, 1, 1), "A", "D-2", "Channel Head", (9.0,), min_t=7.5)]
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "rejected")
        self.assertTrue(any("keeps one value" in e for e in res["errors"]), res["errors"])
        self.assertIsNone(self.equipment("051-V-11").dish_min_thickness)    # a rejected file changes nothing
        self.assertEqual(self.db.query(InspectionRecord).count(), 0)

    def test_each_zone_type_uses_its_own_minimum(self):
        eq = self.equipment("051-V-11")
        eq.min_required_thickness, eq.dish_min_thickness, eq.tube_min_thickness, eq.bottom_min_thickness = D("8"), D("7"), D("1"), D("4")
        eq.roof_min_thickness = D("3")
        self.db.commit()
        rows = [row("051-V-11", Dt(2026, 1, 1), "A", code, typ, (20.0,)) for code, typ in
                (("S-1", "Shell"), ("D-1", "Dish"), ("CH-1", "Channel Head"), ("T-1", "Tube"), ("B-1", "Tank Bottom"), ("R-1", "Roof"))]
        res = self.run_import(make_xlsx(rows))
        self.assertEqual(res["status"], "success", res["errors"])
        got = {z["code"]: D(z["t_min"]) for z in res["summary"][0]["zones"]}
        self.assertEqual(got, {"S-1": D("8"), "D-1": D("7"), "CH-1": D("7"), "T-1": D("1"), "B-1": D("4"), "R-1": D("3")})

    def test_existing_inspection_on_the_same_date_is_replaced(self):
        self.run_import(make_xlsx([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (10.0, 9.0))]))
        res = self.run_import(make_xlsx([row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (10.0, 9.5))]))
        self.assertEqual(res["records_updated"], 1)
        recs = self.records("51-E-01")
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0].min_measured_thickness, D("9.5"))

    # ------------------------------------------------------------------ the downloadable template
    def test_downloaded_template_round_trip(self):
        data = tpl.build_template_workbook()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Instructions", "Inspection Data", "Example"])
        self.assertEqual([c.value for c in wb["Inspection Data"][1]], HEADER)
        # empty template -> clear message
        res = self.run_import(data)
        self.assertEqual(res["status"], "rejected")
        self.assertTrue(any("no data rows" in e for e in res["errors"]))
        # type the rows of the client's file into the template's own data sheet
        ws = wb["Inspection Data"]
        for r in [row("51-E-01", Dt(2024, 1, 1), "John", "S-1", "Shell", (10.01, 10.05, 10.06, 10.08)),
                  row("51-E-01", Dt(2024, 1, 1), "John", "S-2", "Shell", (10.04, 10.04, 10.0, 10.02))]:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        res = self.run_import(buf.getvalue())
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(self.records("51-E-01")[0].min_measured_thickness, D("10.0"))

    def test_example_sheet_is_never_imported(self):
        data = tpl.build_template_workbook()
        self.assertEqual(self.run_import(data)["records_created"], 0)
        self.assertEqual(self.db.query(InspectionRecord).count(), 0)

    def test_damaged_header_of_the_template_sheet(self):
        data = make_xlsx([["51-E-01", Dt(2026, 1, 1)]], header=["a", "b"])
        with self.assertRaises(ValueError):
            self.run_import(data)

    def test_csv_template(self):
        rows = [row("51-E-01", Dt(2026, 1, 1), "A", "S-1", "Shell", (9.9, 9.5, 9.8)),
                row("51-E-01", Dt(2026, 1, 1), "A", "S-2", "Shell", (9.7, 9.6))]
        res = self.run_import(make_csv(rows), "t.csv")
        self.assertEqual(res["status"], "success", res["errors"])
        self.assertEqual(self.records("51-E-01")[0].min_measured_thickness, D("9.5"))

    # ------------------------------------------------------------------ the older flat layout no longer guesses
    def test_old_flat_layout_reads_every_zone_row(self):
        header = ["Equipment Tag", "Inspection Date", "TML Code", "TML Type", "North", "South", "East", "West"]
        data = make_xlsx([["51-E-01", Dt(2026, 1, 1), "S-1", "Shell", 10.0, 10.1, 10.2, 10.3],
                          ["51-E-01", Dt(2026, 1, 1), "S-2", "Shell", 9.7, 9.8, 9.9, 10.0]], header=header, sheet="Sheet1")
        res = self.run_import(data)
        self.assertEqual(res["records_created"], 1, res)
        self.assertEqual(self.db.query(InspectionTML).count(), 2)
        self.assertEqual(self.records("51-E-01")[0].min_measured_thickness, D("9.7"))
        self.assertTrue(res["warnings"])

    def test_old_flat_layout_without_type_is_refused_not_defaulted(self):
        header = ["Equipment Tag", "Inspection Date", "TML Code", "North", "South", "East", "West"]
        data = make_xlsx([["51-E-01", Dt(2026, 1, 1), "T-1", 2.0, 2.1, 2.2, 2.3]], header=header, sheet="Sheet1")
        res = self.run_import(data)
        self.assertEqual(res["records_created"], 0)
        self.assertTrue(any("no valid TML Type" in e for e in res["errors"]), res["errors"])


if __name__ == "__main__":
    unittest.main()
