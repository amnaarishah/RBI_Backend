"""Excel import of the client's Equipment Register Rev02 (wide layout) on a TEMPORARY database.

Expected readings are read straight from the sheet cells, expected results are computed with the
sheet's formulas (short / long term CR, MAX, remaining life, tank bottom x2). The t_min values
used below are only test inputs for the Equipment Profile fields.

Run from backend/:  python -m unittest tests.test_import_rev02 -v
"""
import os
import sys
import unittest
import uuid
from datetime import datetime
from decimal import Decimal as D

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl  # noqa: E402
from sqlalchemy import create_engine, text  # noqa: E402
from sqlalchemy.engine import make_url  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app.core.config import settings  # noqa: E402
from app.core.database import Base  # noqa: E402
import app.models.equipment, app.models.inspection, app.models.cml  # noqa: E402,F401
from app.models.equipment import Equipment  # noqa: E402
from app.models.inspection import InspectionRecord  # noqa: E402
from app.models.cml import InspectionTML, InspectionReading  # noqa: E402
from app.crud import inspection_import as crud  # noqa: E402

REV02 = os.path.join(os.path.dirname(__file__), "..", "..", "Equipment Register-Rev02.xlsx")
BASE_URL = make_url(settings.DATABASE_URL)
TEST_DB = "ams_test_" + uuid.uuid4().hex[:8]
engine = create_engine(BASE_URL.set(database=TEST_DB), pool_pre_ping=True)
Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def col(letter):
    return openpyxl.utils.column_index_from_string(letter)


def block(ws, rows, date_col, zones):
    """{zone: [(date, [readings]) ...]} read from the sheet."""
    out = {}
    for name, cols in zones.items():
        out[name] = [(ws.cell(r, col(date_col)).value,
                      [ws.cell(r, col(c)).value for c in cols]) for r in rows]
    return out


def sheet_life(history, tmin, bottom=False):
    """Controlling remaining life of the latest reading (sheet rows 21-26) -> min over points."""
    (d0, ti), (dp, tp), (df, tf) = history[0], history[-2], history[-1]
    if len(history) == 2:
        dp, tp = d0, ti
    best = None
    for i, cur in enumerate(tf):
        if cur is None:
            continue
        lt = (ti[i] - cur) / ((df - d0).days / 365)
        st = (tp[i] - cur) / ((df - dp).days / 365)
        cr = max(lt, st) if len(history) > 2 else lt
        if cr <= 0:
            continue
        rl = (cur - tmin) / (cr * (2 if bottom else 1))
        rl = max(rl, 0)
        best = rl if best is None else min(best, rl)
    return best


def setUpModule():
    admin = create_engine(BASE_URL.set(database="postgres"), isolation_level="AUTOCOMMIT")
    with admin.connect() as c:
        c.execute(text(f'CREATE DATABASE "{TEST_DB}"'))
    admin.dispose()
    Base.metadata.create_all(bind=engine)


def tearDownModule():
    engine.dispose()
    admin = create_engine(BASE_URL.set(database="postgres"), isolation_level="AUTOCOMMIT")
    with admin.connect() as c:
        c.execute(text(f'DROP DATABASE IF EXISTS "{TEST_DB}" WITH (FORCE)'))
    admin.dispose()


@unittest.skipUnless(os.path.exists(REV02), "Equipment Register-Rev02.xlsx not found")
class Rev02Import(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db = Session()
        eqs = {
            "00-V-001": dict(equipment_type="Vessel", min_required_thickness=D("4.5"), dish_min_thickness=D("6")),
            "00-V-15": dict(equipment_type="Vessel", min_required_thickness=D("9.5"), dish_min_thickness=D("9.5")),
            "00-E-12": dict(equipment_type="Heat Exchanger", min_required_thickness=D("8"),
                            dish_min_thickness=D("8"), tube_min_thickness=D("0.8")),
            "06-TK-12": dict(equipment_type="Tank", min_required_thickness=D("5"), roof_min_thickness=D("4"),
                             bottom_min_thickness=D("2.54")),
        }
        for tag, kw in eqs.items():
            cls.db.add(Equipment(equipment_id=tag, equipment_name=tag, plant="P", unit="U", **kw))
        cls.db.commit()
        cls.data = open(REV02, "rb").read()
        cls.result = crud.import_inspections_from_excel(cls.db, cls.data, "Equipment Register-Rev02.xlsx")

        wb = openpyxl.load_workbook(REV02, data_only=True)
        main, tank = wb["List of equipment"], wb["Tank"]
        four = lambda *starts: {n: [c for c in cols] for n, cols in starts}
        std = {"S1": ["AZ", "BA", "BB", "BC"], "S2": ["BE", "BF", "BG", "BH"],
               "D1": ["BJ", "BK", "BL", "BM"], "D2": ["BO", "BP", "BQ", "BR"]}
        cls.sheet = {
            "00-V-001": block(main, [5, 6], "BV", std),
            "00-V-15": block(main, [18, 19, 20], "BV", std),
            "00-E-12": block(main, [52, 53, 54], "BX", {**std, "TUBE": ["BT"]}),
            "06-TK-12": block(tank, [2, 3, 4], "AC", {
                "C1": ["C", "D", "E", "F"], "C2": ["H", "I", "J", "K"], "C3": ["M", "N", "O", "P"],
                "ROOF": ["R", "S", "T", "U"], "BOTTOM": ["W", "X", "Y"]}),
        }

    @classmethod
    def tearDownClass(cls):
        cls.db.close()

    def zones(self, tag):
        eq = self.db.query(Equipment).filter_by(equipment_id=tag).one()
        return self.db.query(InspectionTML).filter_by(equipment_id=eq.id).order_by(InspectionTML.sort_order).all()

    def test_import_reports_no_errors_and_expected_counts(self):
        self.assertEqual(self.result["errors"], [])
        # 2 + 3 + 3 inspection dates in the register, 3 in the Tank sheet
        self.assertEqual(self.result["records_created"], 11)
        self.assertEqual({t: len(self.zones(t)) for t in self.sheet},
                         {"00-V-001": 4, "00-V-15": 4, "00-E-12": 5, "06-TK-12": 5})

    def test_every_reading_equals_the_sheet(self):
        for tag, zones in self.sheet.items():
            imported = self.zones(tag)
            self.assertEqual(len(imported), len(zones), tag)
            for tml, (zname, history) in zip(imported, zones.items()):
                got = [(r.inspection_date.replace(tzinfo=None),
                        [None if v is None else float(v) for v in
                         (r.reading_1, r.reading_2, r.reading_3, r.reading_4, r.reading_5, r.reading_6)])
                       for r in tml.readings]
                self.assertEqual(len(got), len(history), f"{tag} {tml.tml_code}")
                for (gd, gv), (sd, sv) in zip(got, history):
                    self.assertEqual(gd.date(), sd.date(), f"{tag} {tml.tml_code}")
                    for i, v in enumerate(sv):
                        self.assertAlmostEqual(gv[i], v, places=4, msg=f"{tag} {tml.tml_code} point {i}")
                    self.assertTrue(all(x is None for x in gv[len(sv):]), f"{tag} {tml.tml_code}")

    def test_zone_types_labels_and_tmin(self):
        v1 = {t.tml_code: t for t in self.zones("00-V-001")}
        self.assertEqual((v1["D-1"].tml_type, v1["D-2"].tml_type), ("Dish", "Dish"))   # "D-1" is a dished end
        self.assertEqual(float(v1["D-1"].min_required_thickness), 6.0)                 # dish t_min, not shell
        self.assertEqual(float(v1["S-1"].min_required_thickness), 4.5)
        self.assertEqual(v1["S-1"].direction_labels, "N,S,E,W")                        # labels come from the sheet
        v15 = self.zones("00-V-15")
        self.assertEqual(v15[0].direction_labels, "0°,90°,180°,270°")
        hx = {t.tml_type: t for t in self.zones("00-E-12")}
        self.assertEqual(float(hx["Channel Head"].min_required_thickness), 8.0)
        self.assertEqual(float(hx["Tube"].min_required_thickness), 0.8)
        tk = {t.tml_type: t for t in self.zones("06-TK-12")}
        self.assertEqual(float(tk["Roof"].min_required_thickness), 4.0)
        self.assertEqual(float(tk["Bottom"].min_required_thickness), 2.54)

    def test_remaining_life_equals_sheet_formulas(self):
        tmins = {"00-V-001": {"Shell": 4.5, "Dish": 6.0}, "00-V-15": {"Shell": 9.5, "Dish": 9.5},
                 "00-E-12": {"Shell": 8.0, "Channel Head": 8.0, "Tube": 0.8},
                 "06-TK-12": {"Shell": 5.0, "Roof": 4.0, "Bottom": 2.54}}
        for tag, zones in self.sheet.items():
            for tml, (zname, history) in zip(self.zones(tag), zones.items()):
                last = tml.readings[-1]
                exp = sheet_life(history, tmins[tag][tml.tml_type], bottom=(tml.tml_type == "Bottom"))
                if exp is None:
                    self.assertIsNone(last.controlling_rl, f"{tag} {tml.tml_code}")
                else:
                    self.assertAlmostEqual(float(last.controlling_rl), exp, delta=max(0.01, exp * 1e-4),
                                           msg=f"{tag} {tml.tml_code} ({tml.tml_type})")

    def test_known_values_from_the_sheet(self):
        by = lambda tag: {t.tml_code: t for t in self.zones(tag)}
        self.assertAlmostEqual(float(by("00-V-15")["Shell-1"].readings[-1].controlling_rl), 6.80, places=2)
        self.assertEqual(by("00-E-12")["Tube"].readings[-1].remaining_life_csv.split(",")[0], "47.48")
        self.assertAlmostEqual(float(by("06-TK-12")["Bottom"].readings[-1].controlling_rl), 25.88, places=2)

    def test_importing_again_does_not_duplicate(self):
        before = (self.db.query(InspectionRecord).count(), self.db.query(InspectionTML).count(),
                  self.db.query(InspectionReading).count())
        res = crud.import_inspections_from_excel(self.db, self.data, "Equipment Register-Rev02.xlsx")
        after = (self.db.query(InspectionRecord).count(), self.db.query(InspectionTML).count(),
                 self.db.query(InspectionReading).count())
        self.assertEqual(before, after)
        self.assertEqual(res["errors"], [])


if __name__ == "__main__":
    unittest.main()
