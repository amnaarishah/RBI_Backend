"""Tank t_min calculator = the client's `tank calculations.xlsx`.

The 5 cases below are the sheet's own rows 5-9 (T-101: D 20 m, H 12 m, E 1, bottom leak protection yes); the expected
values are the sheet's cached results (row 5 has no formula in column Q: its 4.32 is the sheet's column-Q formula applied).
Run from backend/:  python -m unittest tests.test_tank_calculator -v
"""
import os
import sys
import unittest
from decimal import Decimal as D

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from fastapi.testclient import TestClient  # noqa: E402
from pydantic import ValidationError  # noqa: E402
from sqlalchemy import func  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.crud.tank_calculator import annular_plate_min, calculate_tank_tmin  # noqa: E402
from app.main import app  # noqa: E402
from app.models.tank import TankMaterial  # noqa: E402
from app.schemas.tank import TankCalculateRequest  # noqa: E402
from app.seed.extract_tank_sheet import read_materials  # noqa: E402
from app.seed.seed_tank_materials import load_json  # noqa: E402

SHEET = os.path.join(os.path.dirname(__file__), "..", "..", "tank calculations.xlsx")

# material, G, S, N (bottom shell), O, P (critical zone), Q (annular), R (roof)
SHEET_ROWS = [
    ("A285 Grade C", "0.9", 163, "6.3309", "1.27", "3.1655", "4.32", "2.2"),
    ("A131 Grade B", "1.5", 172, "9.9994", "1.27", "4.9997", "6", "2.2"),
    ("A573 Grade 450", "0.85", 192, "5.0761", "1.27", "3", "5.84", "2.2"),
    ("A573 Grade 400", "0.85", 172, "5.6663", "1.27", "3", "5.08", "2.2"),
    ("Unknown Grade", "1.1", 163, "7.7378", "1.27", "3.8689", "6", "2.2"),
]


def request(material, g, d_m=20, h_m=12, e="1", leak=True):
    return TankCalculateRequest(material=material, specific_gravity=D(g), diameter_mm=D(d_m) * 1000,
                                height_mm=D(h_m) * 1000, joint_efficiency=D(e), bottom_leak_protection=leak)


class AnnularTable(unittest.TestCase):
    """Column Q of the sheet: G < 1 / G >= 1, band of the shell thickness N, band of S (strictly below the limit)."""

    def test_low_gravity_bands(self):
        a = lambda n, s: annular_plate_min(D("0.9"), D(n), D(s))
        self.assertEqual(a("10", "167.99"), D("4.32"))
        self.assertEqual(a("10", "168"), D("5.08"))           # 168 is not below 168
        self.assertEqual(a("10", "204.99"), D("5.84"))
        self.assertEqual(a("10", "222.99"), D("7.62"))
        self.assertIsNone(a("10", "223"))                     # Out of Range
        self.assertEqual(a("19", "100"), D("4.32"))           # N <= 19
        self.assertEqual(a("19.0001", "100"), D("4.32"))      # 25 band, same value in the first column
        self.assertEqual(a("25.5", "180"), D("6.6"))          # N <= 32: S 168..186 -> 6.6
        self.assertEqual(a("25.5", "190"), D("9.65"))         # N <= 32: S 186..205 -> 9.65
        self.assertEqual(a("38", "100"), D("5.59"))
        self.assertEqual(a("38.0001", "100"), D("6.86"))      # above 38: last row, however thick
        self.assertEqual(a("90", "200"), D("13.46"))

    def test_high_gravity_bands(self):
        a = lambda n, s: annular_plate_min(D("1"), D(n), D(s))   # G = 1 belongs to the G >= 1 table
        self.assertEqual(a("10", "189.99"), D("6"))
        self.assertEqual(a("10", "219.99"), D("7"))
        self.assertEqual(a("10", "249.99"), D("9"))
        self.assertIsNone(a("10", "250"))
        self.assertEqual(a("40", "100"), D("8"))
        self.assertEqual(a("45", "249"), D("19"))
        self.assertIsNone(a("45.0001", "100"))                # thicker than the last row: Out of Range

    def test_the_gravity_switch(self):
        self.assertEqual(annular_plate_min(D("0.9999"), D("10"), D("100")), D("4.32"))
        self.assertEqual(annular_plate_min(D("1"), D("10"), D("100")), D("6"))


class Calculator(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db = SessionLocal()
        if cls.db.query(func.count(TankMaterial.id)).scalar() == 0:
            raise unittest.SkipTest("tank_material is empty (run: python -m app.seed.seed_tank_materials)")

    @classmethod
    def tearDownClass(cls):
        cls.db.close()

    def test_the_five_rows_of_the_sheet(self):
        for mat, g, s, n, o, p, q, r in SHEET_ROWS:
            res = calculate_tank_tmin(self.db, request(mat, g))
            self.assertEqual(res.allowable_stress_mpa, D(s), mat)
            self.assertEqual(res.t_shell_bottom_course_mm, D(n), mat)
            self.assertEqual(res.t_bottom_mm, D(o), mat)
            self.assertEqual(res.t_bottom_critical_zone_mm, D(p), mat)
            self.assertEqual(res.t_annular_mm, D(q), mat)
            self.assertEqual(res.t_roof_mm, D(r), mat)

    def test_bottom_without_leak_protection_is_2_54(self):
        self.assertEqual(calculate_tank_tmin(self.db, request("A285 Grade C", "0.9", leak=False)).t_bottom_mm, D("2.54"))

    def test_critical_zone_is_at_least_3(self):
        r = calculate_tank_tmin(self.db, request("A285 Grade C", "0.9", d_m=5, h_m=4))          # small tank: N is tiny
        self.assertLess(r.t_shell_bottom_course_mm, 6)
        self.assertEqual(r.t_bottom_critical_zone_mm, D("3"))

    def test_joint_efficiency_divides(self):
        a = calculate_tank_tmin(self.db, request("A285 Grade C", "0.9", e="1")).t_shell_bottom_course_mm
        b = calculate_tank_tmin(self.db, request("A285 Grade C", "0.9", e="0.85")).t_shell_bottom_course_mm
        self.assertAlmostEqual(float(b), float(a) / 0.85, places=3)

    def test_millimetres_in_metres_out(self):
        r = calculate_tank_tmin(self.db, request("A285 Grade C", "0.9"))
        self.assertEqual((r.diameter_m, r.height_m), (D("20"), D("12")))

    def test_unknown_material_is_explicit(self):
        r = calculate_tank_tmin(self.db, request("Unknown Grade", "1.1"))
        self.assertTrue(r.is_unknown_material)
        self.assertTrue(any("Unknown material" in w for w in r.warnings))
        self.assertFalse(calculate_tank_tmin(self.db, request("A285 Grade C", "0.9")).warnings)

    def test_stress_above_the_note_of_cell_J1_warns(self):
        r = calculate_tank_tmin(self.db, request("A537 Grade 2", "1.5"))       # S = 236 MPa
        self.assertTrue(any("222" in w for w in r.warnings))

    def test_out_of_range_annular_gives_no_value(self):
        r = calculate_tank_tmin(self.db, request("A537 Grade 2", "0.9"))       # G < 1 table stops below 223 MPa
        self.assertIsNone(r.t_annular_mm)
        self.assertTrue(any("Out of Range" in w for w in r.warnings))
        self.assertGreater(r.t_shell_bottom_course_mm, 0)                      # the other values are still given

    def test_refusals(self):
        with self.assertRaises(ValueError):
            calculate_tank_tmin(self.db, request("A999 Grade Z", "0.9"))                 # not in the list: never guessed
        with self.assertRaises(ValueError):
            calculate_tank_tmin(self.db, request("A285 Grade C", "0.9", h_m="0.3"))      # H must exceed 0.3 m
        for bad in ({"specific_gravity": "0"}, {"diameter_mm": "-1"}, {"joint_efficiency": "1.2"}, {"joint_efficiency": "0"}):
            base = dict(material="A285 Grade C", specific_gravity="0.9", diameter_mm="20000", height_mm="12000",
                        joint_efficiency="1", bottom_leak_protection=True)
            base.update(bad)
            with self.assertRaises(ValidationError):
                TankCalculateRequest(**base)


class Api(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    def test_materials_and_calculate(self):
        mats = self.client.get("/api/v1/tank/materials")
        self.assertEqual(mats.status_code, 200)
        names = [m["name"] for m in mats.json()]
        self.assertIn("A285 Grade C", names)
        self.assertEqual(len(names), len(set(names)))
        res = self.client.post("/api/v1/tank/calculate", json={
            "material": "A285 Grade C", "specific_gravity": "0.9", "diameter_mm": "20000", "height_mm": "12000",
            "joint_efficiency": "1", "bottom_leak_protection": True})
        self.assertEqual(res.status_code, 200, res.text)
        body = res.json()
        self.assertAlmostEqual(float(body["t_shell_bottom_course_mm"]), 6.3309, places=4)
        self.assertAlmostEqual(float(body["t_annular_mm"]), 4.32, places=2)

    def test_unknown_material_is_422(self):
        res = self.client.post("/api/v1/tank/calculate", json={
            "material": "nonsense", "specific_gravity": "0.9", "diameter_mm": "20000", "height_mm": "12000",
            "joint_efficiency": "1", "bottom_leak_protection": True})
        self.assertEqual(res.status_code, 422)


class ProfileFields(unittest.TestCase):
    """The two new result fields are stored on the equipment profile (temporary database; nothing else changes)."""

    def test_new_fields_round_trip(self):
        import uuid
        from sqlalchemy import create_engine, text
        from sqlalchemy.engine import make_url
        from sqlalchemy.orm import sessionmaker
        from app.core.config import settings
        from app.core.database import Base
        import app.models.equipment, app.models.inspection, app.models.cml  # noqa: F401
        from app.models.equipment import Equipment
        from app.schemas.equipment import EquipmentCreate, EquipmentResponse

        base = make_url(settings.DATABASE_URL)
        name = "ams_test_" + uuid.uuid4().hex[:8]
        admin = create_engine(base.set(database="postgres"), isolation_level="AUTOCOMMIT")
        with admin.connect() as c:
            c.execute(text(f'CREATE DATABASE "{name}"'))
        engine = create_engine(base.set(database=name))
        try:
            Base.metadata.create_all(engine)
            db = sessionmaker(bind=engine)()
            data = EquipmentCreate(equipment_id="T-1", equipment_name="tank", equipment_type="Storage Tank",
                                   bottom_min_thickness=D("1.27"), bottom_critical_zone_min_thickness=D("3.1655"),
                                   annular_min_thickness=D("4.32"), tank_specific_gravity=D("0.9"), tank_height=D("12000"),
                                   tank_joint_efficiency=D("0.85"), tank_bottom_leak_protection=False)
            eq = Equipment(**data.model_dump())
            db.add(eq)
            db.commit()
            db.refresh(eq)
            self.assertEqual((eq.bottom_critical_zone_min_thickness, eq.annular_min_thickness), (D("3.1655"), D("4.32")))
            resp = EquipmentResponse.model_validate(eq)
            self.assertEqual(resp.annular_min_thickness, D("4.32"))
            # the calculator inputs are kept too (False = "No" is a real answer, not "not answered")
            self.assertEqual((resp.tank_specific_gravity, resp.tank_height, resp.tank_joint_efficiency), (D("0.9"), D("12000"), D("0.85")))
            self.assertIs(resp.tank_bottom_leak_protection, False)
            other = Equipment(**EquipmentCreate(equipment_id="V-1", equipment_name="v", equipment_type="Pressure Vessel").model_dump())
            db.add(other)
            db.commit()
            self.assertIsNone(other.annular_min_thickness)          # equipment without a value is unaffected
            self.assertIsNone(other.tank_bottom_leak_protection)
            db.close()
        finally:
            engine.dispose()
            with admin.connect() as c:
                c.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
            admin.dispose()


@unittest.skipUnless(os.path.exists(SHEET), "tank calculations.xlsx not found")
class SheetAgreement(unittest.TestCase):
    """The stored JSON (and so the database) is exactly the material table of the client's sheet."""

    def test_json_equals_the_sheet(self):
        self.assertEqual(load_json(), read_materials(SHEET))

    def test_database_equals_the_json(self):
        db = SessionLocal()
        try:
            rows = {m.name: m for m in db.query(TankMaterial).all()}
        finally:
            db.close()
        if not rows:
            self.skipTest("tank_material is empty")
        for r in load_json():
            m = rows[r["name"]]
            self.assertEqual(float(m.s_lower_mpa), r["s_lower_mpa"], r["name"])
            self.assertEqual(float(m.yield_mpa), r["yield_mpa"], r["name"])
        self.assertEqual(len(rows), len(load_json()))

    def test_annular_tables_equal_the_links_sheet(self):
        import openpyxl
        ws = openpyxl.load_workbook(SHEET, data_only=True)["links"]
        from app.crud.tank_calculator import ANNULAR_HIGH_G, ANNULAR_LOW_G
        low = [[ws.cell(r, c).value for c in range(3, 7)] for r in range(3, 8)]        # C3:F7
        high = [[ws.cell(r, c).value for c in range(3, 7)] for r in range(12, 17)]      # C12:F16
        self.assertEqual([v for _, v in ANNULAR_LOW_G["rows"]], low)
        self.assertEqual([v for _, v in ANNULAR_HIGH_G["rows"]], high)


if __name__ == "__main__":
    unittest.main()
